import sys
import json
import copy
import re
from dataclasses import dataclass
from typing import List, Dict, Optional, Tuple
from PySide6.QtWidgets import QApplication, QMainWindow, QWidget, QVBoxLayout, QHBoxLayout, QTableWidget, QTableWidgetItem, QListWidget, QListWidgetItem, QPushButton, QFileDialog, QLabel, QSplitter, QMessageBox, QFormLayout, QLineEdit, QGroupBox, QCheckBox, QScrollArea, QTabWidget, QSpinBox, QAbstractItemView, QDialog, QDialogButtonBox, QColorDialog, QMenu, QInputDialog, QTableWidgetSelectionRange, QFrame, QComboBox
from PySide6.QtCore import Qt
from PySide6.QtGui import QColor, QFont, QKeySequence, QShortcut, QBrush, QAction

@dataclass
class CellItem:
    id_form: int
    label: str
    codigo: str
    tipo: int
    deci: int
    posicion: str
    valor: str

def parse_pos(pos: str) -> Tuple[int, int]:
    r, c = pos.split(":")
    return int(r), int(c)

def fmt_pos(r: int, c: int) -> str:
    return f"{r}:{c}"

def col_name(idx: int) -> str:
    s = ""
    x = idx
    while True:
        s = chr(ord('A') + (x % 26)) + s
        x = x // 26 - 1
        if x < 0:
            break
    return s

class GridEditor(QMainWindow):
    def __init__(self):
        super().__init__()
        self.setWindowTitle("Cierres Maker")
        self.items: List[CellItem] = []
        self.items_by_codigo: Dict[str, CellItem] = {}
        self.pos_to_item: Dict[str, CellItem] = {}
        self.current_codigo: Optional[str] = None
        self.groups: Dict[str, Dict[str, CellItem]] = {}
        self.root_data = None
        self.global_ids: Dict[str, Optional[int]] = {"dia": None, "semana": None, "mes": None, "anio": None}
        self.json_periods: List[str] = []  # períodos que realmente existen en el JSON cargado
        self.updating = False
        self.undo_stack = []
        self.redo_stack = []
        self.copied_data: Optional[Dict] = None
        self.user_highlights: Dict[Tuple[str, str], QColor] = {}

        self.tabs = QTabWidget()
        self.tabs.currentChanged.connect(self.on_tab_changed)
        self.tables: Dict[str, QTableWidget] = {}
        self.current_period: str = "dia"
        self.init_table_for_period("dia")

        self.list = QListWidget()
        self.list.currentRowChanged.connect(self.on_list_change)

        self.load_btn = QPushButton("Cargar JSON")
        self.load_btn.clicked.connect(self.on_load_json)

        self.save_btn = QPushButton("Guardar JSON")
        self.save_btn.clicked.connect(self.on_save_json)

        self.add_row_btn = QPushButton("Agregar Fila")
        self.add_row_btn.clicked.connect(self.on_insert_row)
        self.add_col_btn = QPushButton("Agregar Columna")
        self.add_col_btn.clicked.connect(self.on_insert_col)
        
        self.undo_btn = QPushButton("Deshacer")
        self.undo_btn.clicked.connect(self.undo)
        self.redo_btn = QPushButton("Rehacer")
        self.redo_btn.clicked.connect(self.redo)
        
        self.copy_btn = QPushButton("Copiar Celda")
        self.copy_btn.clicked.connect(self.copy_selection)
        self.paste_btn = QPushButton("Pegar Celda")
        self.paste_btn.clicked.connect(self.paste_selection)

        self.copy_period_btn = QPushButton("📋 Copiar a período")
        self.copy_period_btn.clicked.connect(self.copy_period_to_period)

        self.mass_codes_btn = QPushButton("⚙️ Configurar códigos masivos")
        self.mass_codes_btn.clicked.connect(self.mass_configure_codes)
        
        self.clear_btn = QPushButton("Limpiar")
        self.clear_btn.clicked.connect(self.on_clear_all)

        self.import_excel_btn = QPushButton("Importar Excel")
        self.import_excel_btn.clicked.connect(self.on_import_excel)

        QShortcut(QKeySequence("Ctrl+Z"), self, self.undo)
        QShortcut(QKeySequence("Ctrl+Y"), self, self.redo)
        QShortcut(QKeySequence("Ctrl+Shift+Z"), self, self.redo)
        QShortcut(QKeySequence("Ctrl+C"), self, self.copy_selection)
        QShortcut(QKeySequence("Ctrl+V"), self, self.paste_selection)
        self._move_shortcuts = []
        for seq, fn in [
            ("Alt+Up", lambda: self.move_selected_block(-1, 0)),
            ("Alt+Down", lambda: self.move_selected_block(1, 0)),
            ("Alt+Left", lambda: self.move_selected_block(0, -1)),
            ("Alt+Right", lambda: self.move_selected_block(0, 1)),
            ("Ctrl+Alt+Up", lambda: self.move_selected_block(-1, 0)),
            ("Ctrl+Alt+Down", lambda: self.move_selected_block(1, 0)),
            ("Ctrl+Alt+Left", lambda: self.move_selected_block(0, -1)),
            ("Ctrl+Alt+Right", lambda: self.move_selected_block(0, 1)),
        ]:
            sc = QShortcut(QKeySequence(seq), self)
            sc.setContext(Qt.ApplicationShortcut)
            sc.activated.connect(fn)
            self._move_shortcuts.append(sc)

        bar = self.menuBar()
        menu_res = bar.addMenu("Resaltar")
        act_resaltar = QAction("Resaltar celda...", self)
        act_resaltar.triggered.connect(self.highlight_current_cell)
        act_quitar = QAction("Quitar resaltado", self)
        act_quitar.triggered.connect(self.clear_highlight_current_cell)
        menu_res.addAction(act_resaltar)
        menu_res.addAction(act_quitar)
        
        menu_mov = bar.addMenu("Mover")
        act_mov = QAction("Mover bloque...", self)
        act_mov.triggered.connect(self.move_selected_block_prompt)
        menu_mov.addAction(act_mov)

        self.move_mode = QCheckBox("Mover grupo con clic")
        self.move_mode.setChecked(False)

        self.current_label = QLabel("Sin selección")

        self.ids_box = QGroupBox("IDs de formularios")
        self.id_form_dia = QLineEdit()
        self.id_form_semana = QLineEdit()
        self.id_form_mes = QLineEdit()
        self.id_form_anio = QLineEdit()
        ids_form = QFormLayout()
        ids_form.addRow("Día", self.id_form_dia)
        ids_form.addRow("Semana", self.id_form_semana)
        ids_form.addRow("Mes", self.id_form_mes)
        ids_form.addRow("Año", self.id_form_anio)
        self.ids_box.setLayout(ids_form)
        self.update_ids_btn = QPushButton("Actualizar IDs")
        self.update_ids_btn.clicked.connect(self.on_update_ids)
        
        self.config_ids_btn = QPushButton("Configurar IDs Globales")
        self.config_ids_btn.clicked.connect(self.prompt_global_ids)

        self.mass_ids_btn = QPushButton("⚡ Actualizar IDs masivos")
        self.mass_ids_btn.clicked.connect(self.mass_update_ids)

        self.detail_box = QGroupBox("Detalle celda")
        self.det_label = QLineEdit()
        self.det_codigo = QLineEdit()
        self.det_posicion = QLineEdit(); self.det_posicion.setReadOnly(True)
        self.det_id = QLineEdit()
        self.det_tipo = QLineEdit()
        self.det_deci = QLineEdit()
        self.det_valor = QLineEdit()
        
        # Connect editing signals
        self.det_label.editingFinished.connect(self.on_detail_edited)
        self.det_codigo.editingFinished.connect(self.on_detail_edited)
        self.det_id.editingFinished.connect(self.on_detail_edited)
        self.det_tipo.editingFinished.connect(self.on_detail_edited)
        self.det_deci.editingFinished.connect(self.on_detail_edited)
        self.det_valor.editingFinished.connect(self.on_detail_edited)

        self.suggest_code_btn = QPushButton("Sugerir Código")
        self.suggest_code_btn.clicked.connect(self.on_suggest_codigo)

        det_form = QFormLayout()
        det_form.addRow("Label", self.det_label)
        det_form.addRow("Código", self.det_codigo)
        det_form.addRow("", self.suggest_code_btn)
        det_form.addRow("Posición", self.det_posicion)
        det_form.addRow("Id form", self.det_id)
        det_form.addRow("Tipo", self.det_tipo)
        det_form.addRow("Deci", self.det_deci)
        det_form.addRow("Valor", self.det_valor)
        self.detail_box.setLayout(det_form)

        # Textos e identidad visual de botones
        self.load_btn.setText("📁 Cargar JSON")
        self.save_btn.setText("💾 Guardar JSON")
        self.save_btn.setObjectName("btn_success")
        self.import_excel_btn.setText("📥 Importar Excel")
        self.import_excel_btn.setObjectName("btn_import")
        self.add_row_btn.setText("➕ Agregar Fila")
        self.add_row_btn.setObjectName("btn_edit")
        self.add_col_btn.setText("📊 Agregar Columna")
        self.add_col_btn.setObjectName("btn_edit")
        self.copy_btn.setText("📋 Copiar")
        self.copy_btn.setObjectName("btn_secondary")
        self.paste_btn.setText("📌 Pegar")
        self.paste_btn.setObjectName("btn_secondary")
        self.copy_period_btn.setText("📋 Copiar a período")
        self.copy_period_btn.setObjectName("btn_edit")
        self.mass_codes_btn.setText("⚙️ Configurar códigos masivos")
        self.mass_codes_btn.setObjectName("btn_edit")
        self.undo_btn.setText("↩ Deshacer")
        self.undo_btn.setObjectName("btn_secondary")
        self.redo_btn.setText("↪ Rehacer")
        self.redo_btn.setObjectName("btn_secondary")
        self.clear_btn.setText("🗑 Limpiar Todo")
        self.clear_btn.setObjectName("btn_danger")
        self.move_mode.setText("↔ Mover grupo con clic")

        # Header
        header = QWidget()
        header_layout = QHBoxLayout(header)
        header_layout.setContentsMargins(8, 4, 8, 0)
        title_box = QWidget()
        title_layout = QVBoxLayout(title_box)
        title_layout.setContentsMargins(0, 0, 0, 0)
        title_layout.setSpacing(0)
        title_main = QLabel("Cierres Maker")
        title_sub = QLabel("Sistema de Gestión de Inventarios")
        title_main.setFont(QFont("Segoe UI", 14, QFont.Bold))
        title_sub.setFont(QFont("Segoe UI", 9))
        title_layout.addWidget(title_main)
        title_layout.addWidget(title_sub)
        header_layout.addWidget(title_box, 1)
        self.count_label = QLabel("0 Items | Cargados")
        self.count_label.setFont(QFont("Segoe UI", 10, QFont.Bold))
        header_layout.addWidget(self.count_label, 0, Qt.AlignRight)

        # Toolbar (search + quick filters)
        toolbar = QWidget()
        toolbar_layout = QHBoxLayout(toolbar)
        toolbar_layout.setContentsMargins(8, 0, 8, 4)
        
        # Row Height
        rh_lbl = QLabel("Alto:")
        self.row_height_spin = QSpinBox()
        self.row_height_spin.setRange(20, 200)
        self.row_height_spin.setValue(24)
        self.row_height_spin.valueChanged.connect(self.on_row_height_change)
        self.row_height_spin.setFixedWidth(60)
        toolbar_layout.addWidget(rh_lbl)
        toolbar_layout.addWidget(self.row_height_spin)
        
        search_box = QWidget()
        sb_layout = QHBoxLayout(search_box)
        sb_layout.setContentsMargins(8, 2, 8, 2)
        search_icon = QLabel("🔍")
        self.search_entry = QLineEdit()
        self.search_entry.setPlaceholderText("Buscar en datos...")
        self.search_entry.textChanged.connect(self.on_search_changed)
        self.search_entry.returnPressed.connect(self.on_search_return)
        sb_layout.addWidget(search_icon)
        sb_layout.addWidget(self.search_entry)
        toolbar_layout.addWidget(search_box, 1)
        filters_label = QLabel("Filtros:")
        toolbar_layout.addWidget(filters_label, 0)
        for name in ["INVENTARIO", "FRUTO", "INGRESO"]:
            btn = QPushButton(name)
            toolbar_layout.addWidget(btn, 0)

        ids_panel = QWidget()
        ids_layout = QVBoxLayout(ids_panel)
        ids_layout.addWidget(self.ids_box)
        ids_layout.addWidget(self.update_ids_btn)
        ids_layout.addWidget(self.config_ids_btn)
        ids_layout.addWidget(self.mass_ids_btn)

        def sidebar_sep():
            line = QFrame()
            line.setFrameShape(QFrame.HLine)
            line.setObjectName("sidebar_sep")
            return line

        # Left sidebar (actions + list)
        left_controls = QWidget()
        left_controls_layout = QVBoxLayout(left_controls)
        left_controls_layout.setContentsMargins(8, 8, 8, 8)
        left_controls_layout.setSpacing(5)
        # Grupo: Archivo
        left_controls_layout.addWidget(self.load_btn)
        left_controls_layout.addWidget(self.save_btn)
        left_controls_layout.addWidget(self.import_excel_btn)
        left_controls_layout.addWidget(sidebar_sep())
        # Grupo: Edición
        left_controls_layout.addWidget(self.add_row_btn)
        left_controls_layout.addWidget(self.add_col_btn)
        left_controls_layout.addWidget(self.copy_btn)
        left_controls_layout.addWidget(self.paste_btn)
        left_controls_layout.addWidget(self.copy_period_btn)
        left_controls_layout.addWidget(self.mass_codes_btn)
        left_controls_layout.addWidget(sidebar_sep())
        # Grupo: Historial
        left_controls_layout.addWidget(self.undo_btn)
        left_controls_layout.addWidget(self.redo_btn)
        left_controls_layout.addWidget(sidebar_sep())
        # Grupo: Opciones + zona peligrosa
        left_controls_layout.addWidget(self.move_mode)
        left_controls_layout.addWidget(self.clear_btn)
        self.delete_btn = QPushButton("🗑 Borrar selección")
        self.delete_btn.setObjectName("btn_danger")
        self.delete_btn.clicked.connect(self.delete_selection)
        left_controls_layout.addWidget(self.delete_btn)
        left_controls_layout.addWidget(self.current_label)

        left_splitter = QSplitter(Qt.Vertical)
        left_splitter.addWidget(left_controls)
        left_splitter.addWidget(self.list)
        left_splitter.setStretchFactor(0, 0)
        left_splitter.setStretchFactor(1, 1)

        left_scroll = QScrollArea()
        left_scroll.setWidgetResizable(True)
        left_scroll.setWidget(left_splitter)
        left_scroll.setMinimumWidth(220)

        # Right details panel (ids + details)
        right_panel = QWidget()
        right_layout = QVBoxLayout(right_panel)
        right_layout.setContentsMargins(8, 8, 8, 8)
        right_layout.addWidget(ids_panel)
        right_layout.addWidget(self.detail_box)
        right_scroll = QScrollArea()
        right_scroll.setWidgetResizable(True)
        right_scroll.setWidget(right_panel)
        right_scroll.setMinimumWidth(240)

        splitter = QSplitter()
        splitter.addWidget(left_scroll)
        splitter.addWidget(self.tabs)
        splitter.addWidget(right_scroll)
        splitter.setStretchFactor(0, 0)
        splitter.setStretchFactor(1, 1)
        splitter.setStretchFactor(2, 0)

        root = QWidget()
        root_layout = QVBoxLayout(root)
        root_layout.setContentsMargins(0, 0, 0, 0)
        root_layout.setSpacing(0)
        root_layout.addWidget(header)
        root_layout.addWidget(toolbar)
        root_layout.addWidget(splitter, 1)
        self.setCentralWidget(root)
        self.setStyleSheet("""
            /* === BASE === */
            QWidget { background: #0A0E27; color: #E0E0F0; font-family: 'Segoe UI', Arial; font-size: 13px; }
            QMainWindow { background: #0A0E27; }

            /* === SCROLLBARS === */
            QScrollBar:vertical { background: #13172B; width: 7px; border-radius: 3px; }
            QScrollBar::handle:vertical { background: #3D3D5C; border-radius: 3px; min-height: 24px; }
            QScrollBar::handle:vertical:hover { background: #5865F2; }
            QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical { height: 0px; }
            QScrollBar:horizontal { background: #13172B; height: 7px; border-radius: 3px; }
            QScrollBar::handle:horizontal { background: #3D3D5C; border-radius: 3px; min-width: 24px; }
            QScrollBar::handle:horizontal:hover { background: #5865F2; }
            QScrollBar::add-line:horizontal, QScrollBar::sub-line:horizontal { width: 0px; }

            /* === SPLITTER === */
            QSplitter::handle { background: #1E1E2E; }
            QSplitter::handle:hover { background: #5865F2; }

            /* === GROUPBOX === */
            QGroupBox {
                background: #1E1E2E;
                border: 1px solid #2D2D44;
                border-radius: 7px;
                padding: 12px 8px 8px 8px;
                margin-top: 14px;
                font-weight: bold;
                font-size: 10px;
                color: #7B8EC8;
                letter-spacing: 1px;
            }
            QGroupBox::title { subcontrol-origin: margin; left: 10px; padding: 0 6px; }

            /* === BUTTONS — base === */
            QPushButton {
                background: #5865F2;
                color: #fff;
                border: none;
                border-radius: 6px;
                padding: 8px 12px;
                font-weight: bold;
                font-size: 12px;
            }
            QPushButton:hover { background: #4752C4; }
            QPushButton:pressed { background: #3A42A8; }
            QPushButton:disabled { background: #2D2D44; color: #55556A; }

            /* === BUTTONS — variantes === */
            QPushButton#btn_success { background: #2E7D52; }
            QPushButton#btn_success:hover { background: #256643; }

            QPushButton#btn_import { background: #2A5298; }
            QPushButton#btn_import:hover { background: #1E3D72; }

            QPushButton#btn_edit { background: #37537A; }
            QPushButton#btn_edit:hover { background: #2A3F5E; }

            QPushButton#btn_secondary { background: #2D2D44; color: #A0A0C0; }
            QPushButton#btn_secondary:hover { background: #3D3D5C; color: #fff; }

            QPushButton#btn_danger { background: #7B2020; color: #FFB0B0; }
            QPushButton#btn_danger:hover { background: #A32828; color: #fff; }

            /* === INPUTS === */
            QLineEdit {
                background: #1E2240;
                color: #E0E0F0;
                border: 1px solid #2D2D54;
                border-radius: 5px;
                padding: 6px 8px;
            }
            QLineEdit:focus { border: 1px solid #5865F2; }
            QLineEdit:read-only { background: #13172B; color: #7B8EC8; border-color: #1E2240; }

            /* === SPINBOX === */
            QSpinBox {
                background: #1E2240;
                color: #E0E0F0;
                border: 1px solid #2D2D54;
                border-radius: 5px;
                padding: 4px 6px;
            }
            QSpinBox:focus { border: 1px solid #5865F2; }
            QSpinBox::up-button, QSpinBox::down-button {
                background: #2D2D54;
                border: none;
                width: 16px;
                border-radius: 3px;
            }
            QSpinBox::up-button:hover, QSpinBox::down-button:hover { background: #5865F2; }

            /* === CHECKBOX === */
            QCheckBox { color: #A0A0C0; spacing: 8px; }
            QCheckBox::indicator {
                width: 15px; height: 15px;
                border-radius: 4px;
                border: 2px solid #3D3D5C;
                background: #1E2240;
            }
            QCheckBox::indicator:checked { background: #5865F2; border-color: #5865F2; }
            QCheckBox::indicator:hover { border-color: #5865F2; }

            /* === LIST === */
            QListWidget {
                background: #13172B;
                color: #C8C8E0;
                border: 1px solid #2D2D44;
                border-radius: 6px;
                outline: none;
            }
            QListWidget::item { padding: 5px 8px; border-radius: 3px; }
            QListWidget::item:selected { background: #5865F2; color: #fff; }
            QListWidget::item:hover:!selected { background: #1E2240; }

            /* === TABS === */
            QTabWidget::pane { border: 1px solid #2D2D44; background: #0A0E27; border-radius: 0px; }
            QTabBar::tab {
                background: #13172B;
                color: #7B8EC8;
                padding: 9px 22px;
                border: none;
                font-weight: bold;
                font-size: 12px;
                min-width: 80px;
            }
            QTabBar::tab:selected { background: #5865F2; color: #fff; }
            QTabBar::tab:hover:!selected { background: #1E2240; color: #C0C0E0; }

            /* === HEADERS tabla === */
            QHeaderView::section {
                background: #1E2240;
                color: #7B8EC8;
                border: none;
                border-right: 1px solid #2D2D44;
                border-bottom: 1px solid #2D2D44;
                padding: 5px;
                font-weight: bold;
                font-size: 11px;
            }

            /* === SCROLL AREA === */
            QScrollArea { background: #0A0E27; border: none; }

            /* === MENU BAR === */
            QMenuBar { background: #0A0E27; color: #A0A0C0; border-bottom: 1px solid #1E2240; padding: 2px; }
            QMenuBar::item:selected { background: #1E2240; color: #fff; border-radius: 4px; }
            QMenu { background: #1E1E2E; color: #E0E0F0; border: 1px solid #2D2D44; border-radius: 6px; padding: 4px; }
            QMenu::item { padding: 6px 20px; border-radius: 4px; }
            QMenu::item:selected { background: #5865F2; }

            /* === TOOLTIP === */
            QToolTip { background: #2D2D44; color: #fff; border: 1px solid #5865F2; border-radius: 4px; padding: 4px 8px; }

            /* === LABEL === */
            QLabel { color: #E0E0F0; background: transparent; }

            /* === SEPARATOR === */
            QFrame#sidebar_sep { color: #2D2D44; background: #2D2D44; border: none; max-height: 1px; }
        """)
    
    def period_title(self, p: str) -> str:
        return {"dia": "Día", "semana": "Semana", "mes": "Mes", "anio": "Año"}.get(p, p.capitalize())
    
    def init_table_for_period(self, period: str):
        if period in self.tables:
            return
        tbl = QTableWidget()
        tbl.setSelectionBehavior(QTableWidget.SelectItems)
        tbl.setSelectionMode(QTableWidget.ExtendedSelection)
        tbl.setFont(QFont("Arial", 11))
        tbl.setAlternatingRowColors(True)
        tbl.setStyleSheet(
            "QTableWidget{background:#0D1128; alternate-background-color:#111630; color:#D8D8F0;"
            " gridline-color:#1E2240; selection-background-color:rgba(88,101,242,70);"
            " selection-color:#fff; border:none; outline:none;}"
            "QTableWidget::item{padding:3px;}"
            "QTableWidget::item:selected{background:rgba(88,101,242,70); color:#fff; border:1px solid #5865F2;}"
            "QTableWidget::item:focus{border:1px solid #5865F2; background:rgba(88,101,242,50);}"
        )
        tbl.setProperty("period", period)
        tbl.cellClicked.connect(lambda r, c, p=period: self.on_cell_clicked_tab(p, r, c))
        tbl.cellChanged.connect(lambda r, c, p=period: self.on_cell_changed_tab(p, r, c))
        tbl.setContextMenuPolicy(Qt.CustomContextMenu)
        tbl.customContextMenuRequested.connect(lambda pos, p=period, t=tbl: self.on_table_context_menu(p, t, pos))
        self.tables[period] = tbl
        self.tabs.addTab(tbl, self.period_title(period))
        # Point helpers to the active table
        self.table = self.tables[self.current_period]
    
    def setup_tabs_from_items(self):
        # Solo crear tabs para períodos que realmente existen en el JSON
        active = self._active_periods()
        order = active if active else ["dia"]

        # Crear tablas solo para períodos activos
        for p in order:
            self.init_table_for_period(p)

        # Eliminar tabs de períodos que ya no están activos
        to_remove = [p for p in list(self.tables.keys()) if p not in order]
        for p in to_remove:
            for i in range(self.tabs.count()):
                w = self.tabs.widget(i)
                if w and w.property("period") == p:
                    self.tabs.removeTab(i)
                    break

        # Verificar si el orden ya es correcto
        if self.tabs.count() == len(order):
            correct = True
            for i, p in enumerate(order):
                w = self.tabs.widget(i)
                if not w or w.property("period") != p:
                    correct = False
                    break
            if correct:
                return

        # Reconstruir orden de tabs sin destruir widgets
        current_p = self.current_period

        while self.tabs.count() > 0:
            self.tabs.removeTab(0)

        for p in order:
            if p in self.tables:
                self.tabs.addTab(self.tables[p], self.period_title(p))

        # Restaurar tab activo (si el período activo ya no existe, usar el primero)
        restored = False
        for i in range(self.tabs.count()):
            if self.tabs.widget(i).property("period") == current_p:
                self.tabs.setCurrentIndex(i)
                restored = True
                break
        if not restored and self.tabs.count() > 0:
            self.tabs.setCurrentIndex(0)
            self.current_period = self.tabs.widget(0).property("period")
            self.table = self.tables[self.current_period]
    
    def on_tab_changed(self, idx: int):
        if idx < 0:
            return
        w = self.tabs.widget(idx)
        if not w:
            return
        period = w.property("period")
        self.current_period = period
        self.table = self.tables[period]
        
        # Update search results for the new period
        if hasattr(self, "search_entry"):
            self.on_search_changed(self.search_entry.text())
        
        self.update_duplicates()
    
    def on_cell_clicked_tab(self, period: str, r0: int, c0: int):
        self.current_period = period
        self.table = self.tables[period]
        self.on_cell_clicked(r0, c0)
    
    def on_cell_changed_tab(self, period: str, r: int, c: int):
        self.current_period = period
        self.table = self.tables[period]
        self.on_cell_changed(r, c)

    def save_state(self):
        state = {
            'items': copy.deepcopy(self.items),
            'global_ids': copy.deepcopy(self.global_ids)
        }
        self.undo_stack.append(state)
        if len(self.undo_stack) > 50:
            self.undo_stack.pop(0)
        self.redo_stack.clear()

    def restore_state(self, state):
        self.items = copy.deepcopy(state['items'])
        self.global_ids = copy.deepcopy(state['global_ids'])
        self.items_by_codigo = {d.codigo: d for d in self.items}
        
        self.pos_to_item = {}
        for d in self.items:
            p = self.get_period(d) or "dia"
            self.pos_to_item[(d.posicion, p)] = d
            
        self.build_groups()
        self.refresh_list()
        self.render_from_items()
        self.update_duplicates()
        
        # Restore selection if possible
        if self.current_codigo and self.current_codigo in self.items_by_codigo:
            item = self.items_by_codigo[self.current_codigo]
            r, c = parse_pos(item.posicion)
            p = self.get_period(item) or "dia"
            if p in self.tables:
                idx = [self.tables[t].property("period") for t in self.tables].index(p) if p in self.tables else 0
                # ensure tab is active
                for i in range(self.tabs.count()):
                    w = self.tabs.widget(i)
                    if w.property("period") == p:
                        self.tabs.setCurrentIndex(i)
                        break
                self.table = self.tables[p]
                self.table.setCurrentCell(r, c)
            self.show_cell_details(r, c)
            self.fill_id_fields(item)
        else:
            # If current selection is gone, try to keep table selection or clear details
            r = self.table.currentRow()
            c = self.table.currentColumn()
            if r >= 0 and c >= 0:
                self.show_cell_details(r, c)
            else:
                self.fill_id_fields(None)

    def undo(self):
        if not self.undo_stack:
            return
        
        current_state = {
            'items': copy.deepcopy(self.items),
            'global_ids': copy.deepcopy(self.global_ids)
        }
        self.redo_stack.append(current_state)
        
        state = self.undo_stack.pop()
        self.restore_state(state)

    def redo(self):
        if not self.redo_stack:
            return
            
        current_state = {
            'items': copy.deepcopy(self.items),
            'global_ids': copy.deepcopy(self.global_ids)
        }
        self.undo_stack.append(current_state)
        
        state = self.redo_stack.pop()
        self.restore_state(state)

    def on_list_change(self, row: int):
        # row param corresponds to visual row in QListWidget, which matches 
        # index in self.items ONLY if not filtered.
        # Now we use UserRole to get the absolute index in self.items
        
        item_widget = self.list.currentItem()
        if not item_widget:
            self.current_codigo = None
            self.current_label.setText("Sin selección")
            self.fill_id_fields(None)
            return
            
        idx = item_widget.data(Qt.UserRole)
        if idx is None or not isinstance(idx, int) or idx < 0 or idx >= len(self.items):
             return

        item = self.items[idx]
        self.current_codigo = item.codigo
        r, c = parse_pos(item.posicion)
        p = self.get_period(item) or "dia"
        # Switch to the item's period tab
        for i in range(self.tabs.count()):
            w = self.tabs.widget(i)
            if w.property("period") == p:
                self.tabs.setCurrentIndex(i)
                break
        
        QApplication.processEvents()

        tbl = self.tables.get(p)
        if tbl and not tbl.item(r, c):
            tbl.setItem(r, c, QTableWidgetItem(""))
        
        if tbl:
            # Force selection and scroll to item
            qitem = tbl.item(r, c)
            if qitem:
                tbl.setCurrentItem(qitem)
                tbl.scrollToItem(qitem, QAbstractItemView.PositionAtCenter)
            else:
                tbl.setCurrentCell(r, c)
                tbl.scrollToItem(tbl.item(r, c), QAbstractItemView.PositionAtCenter)
            
            self.table = tbl
            tbl.setFocus()
            
        self.show_cell_details(r, c)
        self.current_label.setText(f"Seleccionado: {item.codigo} | Fila {r} Col {c}")
        self.fill_id_fields(item)

    def on_cell_clicked(self, r0: int, c0: int):
        r = r0
        c = c0
        self.table.setCurrentCell(r, c)
        self.show_cell_details(r, c)
        if self.current_codigo and self.move_mode.isChecked():
            self.save_state()
            itm = self.items_by_codigo[self.current_codigo]
            self.move_group_for_item(itm, r, c)
            self.update_duplicates()

    def copy_selection(self):
        ranges = self.table.selectedRanges()
        if not ranges:
            self.copied_data = None
            self.current_label.setText("Nada para copiar")
            return
        if len(ranges) > 1:
            self.copied_data = None
            self.current_label.setText("Selecciona un solo bloque para copiar")
            return
        rg = ranges[0]
        top, left, bottom, right = rg.topRow(), rg.leftColumn(), rg.bottomRow(), rg.rightColumn()
        rows = bottom - top + 1
        cols = right - left + 1
        block = []
        for rr in range(rows):
            row_data = []
            for cc in range(cols):
                pos = fmt_pos(top + rr, left + cc)
                item = self.get_item_at(pos, self.current_period)
                if item:
                    row_data.append({
                        "label": item.label,
                        "codigo": item.codigo,
                        "id_form": item.id_form,
                        "tipo": item.tipo,
                        "deci": item.deci,
                        "valor": item.valor
                    })
                else:
                    row_data.append(None)
            block.append(row_data)
        self.copied_data = {
            "kind": "range",
            "period": self.current_period,
            "rows": rows,
            "cols": cols,
            "block": block
        }
        self.current_label.setText(f"Copiado: {rows}x{cols}")

    def paste_selection(self):
        if not hasattr(self, 'copied_data') or not self.copied_data:
            return
        
        r = self.table.currentRow()
        c = self.table.currentColumn()
        if r < 0 or c < 0:
            return
            
        self.save_state()

        data = self.copied_data
        if isinstance(data, dict) and data.get("kind") == "range":
            rows = int(data.get("rows", 0))
            cols = int(data.get("cols", 0))
            block = data.get("block") or []
            if rows <= 0 or cols <= 0:
                return
            needed_r = r + rows
            needed_c = c + cols
            for tbl in self.tables.values():
                if tbl.rowCount() < needed_r:
                    tbl.setRowCount(needed_r)
                if tbl.columnCount() < needed_c:
                    tbl.setColumnCount(needed_c)

            target_id = self.global_ids.get(self.current_period)
            if target_id is None:
                target_id = 0

            self.updating = True
            try:
                for rr in range(rows):
                    if rr >= len(block):
                        continue
                    for cc in range(cols):
                        if cc >= len(block[rr]):
                            continue
                        cell_data = block[rr][cc]
                        if not cell_data:
                            continue
                        dr = r + rr
                        dc = c + cc
                        pos = fmt_pos(dr, dc)
                        existing = self.pos_to_item.get((pos, self.current_period))
                        new_code = (cell_data.get("codigo") or "").strip().upper()
                        if existing:
                            old_code = (existing.codigo or "").strip().upper()
                            if old_code and old_code in self.items_by_codigo and self.items_by_codigo.get(old_code) == existing and old_code != new_code:
                                del self.items_by_codigo[old_code]
                            existing.label = cell_data.get("label", "")
                            existing.codigo = new_code
                            existing.id_form = int(target_id) if target_id else int(cell_data.get("id_form", 0) or 0)
                            existing.tipo = int(cell_data.get("tipo", 0) or 0)
                            existing.deci = int(cell_data.get("deci", 0) or 0)
                            existing.valor = cell_data.get("valor", "")
                            if existing.codigo:
                                self.items_by_codigo[existing.codigo] = existing
                            item_obj = existing
                        else:
                            item_obj = CellItem(
                                id_form=int(target_id) if target_id else int(cell_data.get("id_form", 0) or 0),
                                label=cell_data.get("label", ""),
                                codigo=new_code,
                                tipo=int(cell_data.get("tipo", 0) or 0),
                                deci=int(cell_data.get("deci", 0) or 0),
                                posicion=pos,
                                valor=cell_data.get("valor", "")
                            )
                            self.items.append(item_obj)
                            self.pos_to_item[(pos, self.current_period)] = item_obj
                            if item_obj.codigo:
                                self.items_by_codigo[item_obj.codigo] = item_obj

                        qitem = self.table.item(dr, dc)
                        if not qitem:
                            qitem = QTableWidgetItem(item_obj.label)
                            self.table.setItem(dr, dc, qitem)
                        else:
                            qitem.setText(item_obj.label)
                        qitem.setData(Qt.UserRole, item_obj.codigo)
            finally:
                self.updating = False

            self.build_groups()
            self.refresh_list()
            self.update_duplicates()
            self.show_cell_details(r, c)
            return

        pos = fmt_pos(r, c)
        item = self.pos_to_item.get((pos, self.current_period))
        if not isinstance(data, dict):
            return
        if "label" not in data and "codigo" not in data:
            return

        if not item:
            item = CellItem(
                id_form=int(data.get("id_form", 0) or 0),
                label=data.get("label", ""),
                codigo=(data.get("codigo") or "").strip().upper(),
                tipo=int(data.get("tipo", 0) or 0),
                deci=int(data.get("deci", 0) or 0),
                posicion=pos,
                valor=data.get("valor", "")
            )
            self.items.append(item)
            self.pos_to_item[(pos, self.current_period)] = item
            if item.codigo:
                self.items_by_codigo[item.codigo] = item
        else:
            old_code = (item.codigo or "").strip().upper()
            new_code = (data.get("codigo") or "").strip().upper()
            if old_code and old_code in self.items_by_codigo and self.items_by_codigo.get(old_code) == item and old_code != new_code:
                del self.items_by_codigo[old_code]
            item.label = data.get("label", "")
            item.codigo = new_code
            item.id_form = int(data.get("id_form", 0) or 0)
            item.tipo = int(data.get("tipo", 0) or 0)
            item.deci = int(data.get("deci", 0) or 0)
            item.valor = data.get("valor", "")
            if item.codigo:
                self.items_by_codigo[item.codigo] = item

        self.place_item(item)
        self.show_cell_details(r, c)
        self.refresh_list()
        self.update_duplicates()

    def delete_selection(self):
        items_to_del = []
        for qitem in self.table.selectedItems():
            r = self.table.row(qitem)
            c = self.table.column(qitem)
            item = self.get_item_at(fmt_pos(r, c), self.current_period)
            if item and item not in items_to_del:
                items_to_del.append(item)
        if not items_to_del:
            return
        self.save_state()
        for item in items_to_del:
            if item in self.items:
                self.items.remove(item)
            if item.codigo and self.items_by_codigo.get(item.codigo) == item:
                del self.items_by_codigo[item.codigo]
            self.pos_to_item.pop((item.posicion, self.current_period), None)
            self.table.setItem(*parse_pos(item.posicion), QTableWidgetItem(""))
        self.show_cell_details(self.table.currentRow(), self.table.currentColumn())
        self.refresh_list()
        self.update_duplicates()
        if hasattr(self, "count_label") and self.count_label:
            self.count_label.setText(f"{len(self.items)} Items | Cargados")

    def update_duplicates(self):
        prev_state = self.updating
        self.updating = True
        try:
            # Contar cuantas veces aparece cada codigo (ignorando vacios)
            tmp_counts: Dict[str, int] = {}
            for d in self.items:
                code = (d.codigo or "").strip().upper()
                if not code:
                    continue
                tmp_counts[code] = tmp_counts.get(code, 0) + 1
            self.duplicate_codes = {k for k, v in tmp_counts.items() if v > 1}

            COLOR_OK        = QColor(0, 0, 0, 0)   # transparente — deja actuar al stylesheet
            COLOR_TEXT_OK   = QColor("#D8D8F0")   # texto claro para tema oscuro
            COLOR_DUP       = QColor("#7A5C00")   # amarillo oscuro para duplicados
            COLOR_TEXT_DUP  = QColor("#FFE580")   # texto dorado sobre fondo oscuro

            # Colorear celda a celda en todas las tablas
            for period, tbl in self.tables.items():
                for r in range(tbl.rowCount()):
                    for c in range(tbl.columnCount()):
                        qitem = tbl.item(r, c)
                        pos = fmt_pos(r, c)
                        
                        # Obtener resaltado manual y datos
                        hl = self.user_highlights.get((pos, period))
                        cell = self.pos_to_item.get((pos, period))
                        
                        # Si hay algo que pintar pero no hay item visual, crearlo
                        if (hl or (cell and cell.codigo)) and not qitem:
                            label = cell.label if cell else ""
                            qitem = QTableWidgetItem(label)
                            tbl.setItem(r, c, qitem)
                            
                        # Si sigue sin haber item (celda vacia sin highlight), saltar
                        if not qitem:
                            continue

                        if cell:
                            code_norm = (cell.codigo or "").strip().upper()
                            if hl:
                                qitem.setBackground(QBrush(hl))
                                qitem.setForeground(QBrush(COLOR_TEXT_OK))
                                qitem.setData(Qt.BackgroundRole, QBrush(hl))
                                qitem.setData(Qt.ForegroundRole, QBrush(COLOR_TEXT_OK))
                                qitem.setToolTip("Resaltado manual")
                            elif code_norm and code_norm in self.duplicate_codes:
                                qitem.setBackground(QBrush(COLOR_DUP))
                                qitem.setForeground(QBrush(COLOR_TEXT_DUP))
                                qitem.setData(Qt.BackgroundRole, QBrush(COLOR_DUP))
                                qitem.setData(Qt.ForegroundRole, QBrush(COLOR_TEXT_DUP))
                                qitem.setToolTip(f"Código duplicado: {cell.codigo}")
                            else:
                                qitem.setBackground(QBrush(COLOR_OK))
                                qitem.setForeground(QBrush(COLOR_TEXT_OK))
                                qitem.setData(Qt.BackgroundRole, QBrush(COLOR_OK))
                                qitem.setData(Qt.ForegroundRole, QBrush(COLOR_TEXT_OK))
                                qitem.setToolTip("")
                        else:
                            # Celda vacia pero con highlight manual
                            if hl:
                                qitem.setBackground(QBrush(hl))
                                qitem.setForeground(QBrush(COLOR_TEXT_OK))
                                qitem.setData(Qt.BackgroundRole, QBrush(hl))
                                qitem.setData(Qt.ForegroundRole, QBrush(COLOR_TEXT_OK))
                                qitem.setToolTip("Resaltado manual")
                            else:
                                qitem.setBackground(QBrush(COLOR_OK))
                                qitem.setForeground(QBrush(COLOR_TEXT_OK))
                                qitem.setData(Qt.BackgroundRole, QBrush(COLOR_OK))
                                qitem.setData(Qt.ForegroundRole, QBrush(COLOR_TEXT_OK))
                                qitem.setToolTip("")
        finally:
            self.updating = prev_state

    def on_row_height_change(self, val: int):
        for tbl in self.tables.values():
            for r in range(tbl.rowCount()):
                tbl.setRowHeight(r, val)

    def highlight_current_cell(self):
        r = self.table.currentRow()
        c = self.table.currentColumn()
        if r < 0 or c < 0:
            return
        color = QColorDialog.getColor(QColor("#FFD35A"), self, "Elegir color")
        if not color.isValid():
            return
        pos = fmt_pos(r, c)
        self.user_highlights[(pos, self.current_period)] = color
        self.update_duplicates()

    def clear_highlight_current_cell(self):
        r = self.table.currentRow()
        c = self.table.currentColumn()
        if r < 0 or c < 0:
            return
        pos = fmt_pos(r, c)
        self.user_highlights.pop((pos, self.current_period), None)
        self.update_duplicates()

    def move_selected_block_prompt(self):
        dr, ok = QInputDialog.getInt(self, "Mover bloque", "Filas (negativo sube):", 0, -10000, 10000, 1)
        if not ok:
            return
        dc, ok = QInputDialog.getInt(self, "Mover bloque", "Columnas (negativo izquierda):", 0, -10000, 10000, 1)
        if not ok:
            return
        if dr == 0 and dc == 0:
            return
        self.move_selected_block(dr, dc)

    def move_selected_block(self, dr: int, dc: int):
        ranges = self.table.selectedRanges()
        if not ranges:
            return
        if len(ranges) > 1:
            QMessageBox.information(self, "Mover", "Selecciona un solo bloque para mover.")
            return
        rg = ranges[0]
        top, left, bottom, right = rg.topRow(), rg.leftColumn(), rg.bottomRow(), rg.rightColumn()
        period = self.current_period

        moving = []
        selected_keys = set()
        for r in range(top, bottom + 1):
            for c in range(left, right + 1):
                pos = fmt_pos(r, c)
                key = (pos, period)
                it = self.pos_to_item.get(key)
                if it:
                    moving.append((r, c, it))
                    selected_keys.add(key)

        if not moving:
            return

        dest_keys = set()
        max_r = 0
        max_c = 0
        for r, c, it in moving:
            nr = r + dr
            nc = c + dc
            if nr < 0 or nc < 0:
                QMessageBox.information(self, "Mover", "El bloque no puede salir del borde superior/izquierdo.")
                return
            npos = fmt_pos(nr, nc)
            nkey = (npos, period)
            dest_keys.add(nkey)
            max_r = max(max_r, nr)
            max_c = max(max_c, nc)

        for nkey in dest_keys:
            if nkey in self.pos_to_item and nkey not in selected_keys:
                QMessageBox.information(self, "Mover", "El bloque choca con celdas ocupadas.")
                return

        for tbl in self.tables.values():
            if tbl.rowCount() <= max_r:
                tbl.setRowCount(max_r + 1)
            if tbl.columnCount() <= max_c:
                tbl.setColumnCount(max_c + 1)

        self.save_state()
        prev = self.updating
        self.updating = True
        # Bloquear señales de la tabla activa durante todo el bloque de movimiento.
        # Esto evita que cellChanged se emita (sincrónica o diferidamente) y
        # dispare la lógica de sync hacia Mes/Año en las nuevas posiciones.
        self.table.blockSignals(True)
        try:
            moved_qitems = {}
            for r, c, it in moving:
                qitem = self.table.takeItem(r, c)
                moved_qitems[(r, c)] = qitem
                old_key = (it.posicion, period)
                self.pos_to_item.pop(old_key, None)

            for r, c, it in moving:
                nr = r + dr
                nc = c + dc
                it.posicion = fmt_pos(nr, nc)
                self.pos_to_item[(it.posicion, period)] = it

                qitem = moved_qitems.get((r, c))
                if not qitem:
                    qitem = QTableWidgetItem(it.label)
                qitem.setText(it.label)
                qitem.setData(Qt.UserRole, it.codigo)
                self.table.setItem(nr, nc, qitem)
        finally:
            self.table.blockSignals(False)
            self.updating = prev

        self.build_groups()
        self.refresh_list()
        self.update_duplicates()

        self.table.clearSelection()
        new_range = QTableWidgetSelectionRange(top + dr, left + dc, bottom + dr, right + dc)
        self.table.setRangeSelected(new_range, True)

    def on_table_context_menu(self, period: str, tbl: QTableWidget, pos):
        idx = tbl.indexAt(pos)
        if idx.isValid():
            tbl.setCurrentCell(idx.row(), idx.column())
        gpos = tbl.viewport().mapToGlobal(pos)
        menu = QMenu(self)
        a1 = QAction("Resaltar celda...", self)
        a2 = QAction("Quitar resaltado", self)
        a3 = QAction("Mover bloque...", self)
        a1.triggered.connect(self.highlight_current_cell)
        a2.triggered.connect(self.clear_highlight_current_cell)
        a3.triggered.connect(self.move_selected_block_prompt)
        menu.addAction(a1)
        menu.addAction(a2)
        menu.addAction(a3)
        menu.exec(gpos)

    def on_search_changed(self, text: str):
        text = text.lower().strip()
        self.list.clear()
        
        # Filter list items
        for idx, d in enumerate(self.items):
            # Only show items belonging to the current active period
            p = self.get_period(d) or "dia"
            if p != self.current_period:
                continue
                
            # Match against code, label, or value
            full_str = f"{d.codigo} | {d.label} | {d.valor}"
            if not text or text in full_str.lower():
                disp = f"{d.codigo} | {d.label}"
                if d.valor:
                    disp += f" | {d.valor}"
                item = QListWidgetItem(disp)
                item.setData(Qt.UserRole, idx)
                self.list.addItem(item)
    
    def on_search_return(self):
        text = self.search_entry.text().lower().strip()
        if not text:
            return
            
        # Find first match in items that is NOT the current selection if possible,
        # or just the first match.
        # For simplicity, let's find the first match in the filtered list.
        
        if self.list.count() > 0:
            # Select first item in the list, which triggers on_list_change -> selecting cell
            self.list.setCurrentRow(0)
            self.list.setFocus()

    def on_load_json(self):
        path, _ = QFileDialog.getOpenFileName(self, "Abrir JSON", "", "Archivos (*.json *.txt);;Todos (*.*)")
        if not path:
            return
        try:
            with open(path, "r", encoding="utf-8") as f:
                raw = f.read()
            try:
                data = json.loads(raw)
            except Exception:
                lines = [json.loads(l) for l in raw.splitlines() if l.strip()]
                data = lines
        except Exception as e:
            QMessageBox.critical(self, "Error", str(e))
            return
        def rows_from_datosAG(x):
            if isinstance(x, dict) and isinstance(x.get("datosAG"), list):
                out = []
                for group in x["datosAG"]:
                    if isinstance(group, list):
                        out.extend([d for d in group if isinstance(d, dict)])
                return out
            return None
        self.root_data = data
        rows = rows_from_datosAG(data)
        if rows is None:
            def collect_entries(x):
                out = []
                def rec(v):
                    if isinstance(v, dict):
                        if all(k in v for k in ("codigo", "posicion", "label")):
                            out.append(v)
                        else:
                            for vv in v.values():
                                rec(vv)
                    elif isinstance(v, list):
                        for e in v:
                            rec(e)
                rec(x)
                return out
            rows = collect_entries(data)
        items: List[CellItem] = []
        for d in rows:
            try:
                items.append(CellItem(
                    id_form=int(d.get("id_form", 0)),
                    label=str(d.get("label", "")),
                    codigo=str(d.get("codigo", "")),
                    tipo=int(d.get("tipo", 0)),
                    deci=int(d.get("deci", 0)),
                    posicion=str(d.get("posicion", "1:1")),
                    valor=str(d.get("valor", "")),
                ))
            except Exception:
                pass
        self.items = items
        self.items_by_codigo = {d.codigo: d for d in self.items}
        
        self.extract_global_ids()
        
        # Rebuild pos_to_item with (pos, period) keys to support overlapping positions
        self.pos_to_item = {}
        for d in self.items:
            p = self.get_period(d) or "dia"
            self.pos_to_item[(d.posicion, p)] = d
            
        self.build_groups()
        
        # Check if any global IDs are missing OR just prompt always as requested
        # User requested: "al inicio quiero que pregunte por el id form... para que se filtre"
        # So we force prompt here.
        self.prompt_global_ids(force_filter=True)
            
        self.refresh_list()
        self.render_from_items()
        self.undo_stack.clear()
        self.redo_stack.clear()
        self.current_label.setText(f"Cargados: {len(self.items)} items")
        if hasattr(self, "count_label") and self.count_label:
            self.count_label.setText(f"{len(self.items)} Items | Cargados")
        if not self.items:
            QMessageBox.information(self, "Aviso", "No se encontraron items válidos en el JSON")

    def prompt_global_ids(self, force_filter=False):
        dlg = QDialog(self)
        dlg.setWindowTitle("Configurar IDs Globales")
        layout = QVBoxLayout(dlg)
        
        form = QFormLayout()
        inputs = {}
        for k in ["dia", "semana", "mes", "anio"]:
            inp = QLineEdit()
            val = self.global_ids.get(k)
            if val is not None:
                inp.setText(str(val))
            inputs[k] = inp
            label = {"dia": "Día", "semana": "Semana", "mes": "Mes", "anio": "Año"}.get(k, k)
            form.addRow(f"ID {label}:", inp)
            
        layout.addLayout(form)
        
        btns = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        btns.accepted.connect(dlg.accept)
        btns.rejected.connect(dlg.reject)
        layout.addWidget(btns)
        
        if dlg.exec() == QDialog.Accepted:
            for k, inp in inputs.items():
                txt = inp.text().strip()
                if txt.isdigit():
                    self.global_ids[k] = int(txt)
                else:
                    self.global_ids[k] = None
            
            if force_filter:
                # Filter items to only keep those matching the configured IDs
                filtered_items = []
                # Collect valid IDs
                valid_ids = {v for v in self.global_ids.values() if v is not None}
                
                for d in self.items:
                    # If item's ID matches one of the valid global IDs, keep it.
                    # Or if the item's period is detected via other means, check if that period is allowed.
                    
                    # Strict filtering based on ID match as requested: "filtre cada valor... para que solo salgan los datos de este tipo"
                    if d.id_form in valid_ids:
                        filtered_items.append(d)
                
                self.items = filtered_items
                self.items_by_codigo = {d.codigo: d for d in self.items}
                
                self.pos_to_item = {}
                for d in self.items:
                    p = self.get_period(d) or "dia"
                    self.pos_to_item[(d.posicion, p)] = d
                    
                self.build_groups()
                QMessageBox.information(self, "Info", f"Datos filtrados. {len(self.items)} items retenidos.")

            self.apply_global_ids_to_root()

    def mass_update_ids(self):
        if not self.items:
            QMessageBox.information(self, "Aviso", "No hay items cargados.")
            return
        dlg = QDialog(self)
        dlg.setWindowTitle("Actualizar IDs masivamente")
        layout = QVBoxLayout(dlg)
        form = QFormLayout()
        inputs = {}
        for k in ["dia", "semana", "mes", "anio"]:
            inp = QLineEdit()
            val = self.global_ids.get(k)
            if val is not None:
                inp.setText(str(val))
            inputs[k] = inp
            label = {"dia": "Día", "semana": "Semana", "mes": "Mes", "anio": "Año"}.get(k, k)
            form.addRow(f"ID {label}:", inp)
        layout.addLayout(form)
        btns = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        btns.accepted.connect(dlg.accept)
        btns.rejected.connect(dlg.reject)
        layout.addWidget(btns)
        if dlg.exec() != QDialog.Accepted:
            return

        new_ids = {}
        for k, inp in inputs.items():
            txt = inp.text().strip()
            new_ids[k] = int(txt) if txt.isdigit() else None

        if not any(v is not None for v in new_ids.values()):
            QMessageBox.warning(self, "Error", "Debe configurar al menos un ID.")
            return

        self.save_state()

        periods = [self.get_period(d) for d in self.items]
        self.global_ids = new_ids

        updated = 0
        skipped = 0
        for d, p in zip(self.items, periods):
            if p is None:
                skipped += 1
                continue
            new_id = new_ids.get(p)
            if new_id is not None:
                d.id_form = new_id
                updated += 1

        self.items_by_codigo = {d.codigo: d for d in self.items}
        self.pos_to_item = {}
        for d in self.items:
            p = self.get_period(d) or "dia"
            self.pos_to_item[(d.posicion, p)] = d
        self.build_groups()
        self.apply_global_ids_to_root()
        self.refresh_list()
        self.update_duplicates()
        self.show_cell_details(self.table.currentRow(), self.table.currentColumn())
        QMessageBox.information(
            self, "Listo",
            f"IDs actualizados: {updated} items.\n"
            f"{skipped} items sin período detectable se dejaron sin tocar."
        )

    def copy_period_to_period(self):
        source_period = self.current_period
        selected = self.table.selectedItems()
        if not selected:
            QMessageBox.information(self, "Aviso", "Selecciona una o varias celdas/columnas primero.")
            return

        targets = [p for p in ["dia", "semana", "mes", "anio"] if p != source_period]
        dlg = QDialog(self)
        dlg.setWindowTitle("Copiar a período")
        layout = QVBoxLayout(dlg)
        form = QFormLayout()
        target_combo = QComboBox()
        for p in targets:
            label = {"dia": "Día", "semana": "Semana", "mes": "Mes", "anio": "Año"}.get(p, p)
            target_combo.addItem(label, p)
        form.addRow("Período destino:", target_combo)
        overwrite_cb = QCheckBox("Sobrescribir celdas existentes")
        overwrite_cb.setChecked(True)
        form.addRow("", overwrite_cb)
        layout.addLayout(form)
        btns = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        btns.accepted.connect(dlg.accept)
        btns.rejected.connect(dlg.reject)
        layout.addWidget(btns)
        if dlg.exec() != QDialog.Accepted:
            return

        target = target_combo.currentData()
        overwrite = overwrite_cb.isChecked()
        target_id = self.global_ids.get(target)
        if target_id is None:
            QMessageBox.warning(
                self, "Error",
                f"El período {self.period_title(target)} no tiene ID configurado.\n"
                "Configúralo en 'Configurar IDs Globales' antes de copiar."
            )
            return

        source_items = []
        for qitem in selected:
            r = self.table.row(qitem)
            c = self.table.column(qitem)
            item = self.get_item_at(fmt_pos(r, c), source_period)
            if item and item not in source_items:
                source_items.append(item)
        if not source_items:
            QMessageBox.information(self, "Aviso", "Las celdas seleccionadas no contienen datos.")
            return

        ret = QMessageBox.question(
            self, "Confirmar",
            f"Copiar {len(source_items)} campo(s) de {self.period_title(source_period)} "
            f"a {self.period_title(target)}?"
        )
        if ret != QMessageBox.Yes:
            return

        self.save_state()

        created = 0
        removed = 0
        skipped = 0
        for src in source_items:
            pos = src.posicion
            existing = self.pos_to_item.get((pos, target))
            if existing:
                if not overwrite:
                    skipped += 1
                    continue
                if existing in self.items:
                    self.items.remove(existing)
                if existing.codigo and self.items_by_codigo.get(existing.codigo) == existing:
                    del self.items_by_codigo[existing.codigo]
                removed += 1
            new = CellItem(
                id_form=target_id,
                label=src.label,
                codigo=self._swap_prefix(src.codigo, target),
                tipo=src.tipo,
                deci=src.deci,
                posicion=pos,
                valor=src.valor,
            )
            self._insert_item(new)
            created += 1

        self.pos_to_item = {}
        for d in self.items:
            p = self.get_period(d) or "dia"
            self.pos_to_item[(d.posicion, p)] = d
        self.items_by_codigo = {d.codigo: d for d in self.items}
        self.build_groups()

        if target not in self.json_periods:
            order = ["dia", "semana", "mes", "anio"]
            self.json_periods = [p for p in order if p in self.json_periods or p == target]

        self.render_from_items()
        self.refresh_list()
        self.update_duplicates()

        msg = f"Copiados {created} campo(s) a {self.period_title(target)}."
        if removed:
            msg += f" Reemplazados {removed}."
        if skipped:
            msg += f" Omitidos {skipped}."
        QMessageBox.information(self, "Listo", msg)

    def mass_configure_codes(self):
        period = self.current_period
        period_items = [d for d in self.items if (self.get_period(d) or "dia") == period]
        if not period_items:
            QMessageBox.information(self, "Aviso", f"No hay items en la hoja {self.period_title(period)}.")
            return

        prefix_map = {"dia": "CD", "semana": "CS", "mes": "CM", "anio": "CA"}
        prefix = prefix_map.get(period, "CD")

        first_code = ""
        for d in period_items:
            if d.codigo:
                first_code = d.codigo
                break

        dlg = QDialog(self)
        dlg.setWindowTitle("Configurar códigos masivos")
        layout = QVBoxLayout(dlg)
        form = QFormLayout()
        inp = QLineEdit()
        inp.setText(first_code)
        inp.selectAll()
        form.addRow(f"Código inicial para {self.period_title(period)}:", inp)
        layout.addLayout(form)
        btns = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        btns.accepted.connect(dlg.accept)
        btns.rejected.connect(dlg.reject)
        layout.addWidget(btns)
        if dlg.exec() != QDialog.Accepted:
            return

        raw = inp.text().strip().upper()
        m = re.match(r'^(CD|CS|CM|CA)_([A-Z]+)?(\d+)([A-Z]*)$', raw)
        if not m:
            QMessageBox.warning(
                self, "Error",
                f"Código inválido. Formato esperado: {prefix}_{texto}{numero}{cliente}, "
                f"ej: {prefix}_MAS000001."
            )
            return
        if m.group(1) != prefix:
            QMessageBox.warning(
                self, "Error",
                f"El código debe comenzar con {prefix} para la hoja {self.period_title(period)}."
            )
            return

        base = m.group(1) + "_" + (m.group(2) or "")
        start_num = int(m.group(3))
        width = len(m.group(3))
        suffix = m.group(4) or ""

        ret = QMessageBox.question(
            self, "Confirmar",
            f"Actualizar {len(period_items)} código(s) de {self.period_title(period)} "
            f"desde {raw} en orden (fila, columna)?"
        )
        if ret != QMessageBox.Yes:
            return

        # Orden (sección de columna, fila, columna): mismo criterio que _rebuild_datosAG
        all_cols = sorted(set(parse_pos(d.posicion)[1] for d in period_items))
        col_to_sec: Dict[int, int] = {}
        if all_cols:
            sec_idx = 0
            col_to_sec[all_cols[0]] = 0
            for i in range(1, len(all_cols)):
                if all_cols[i] - all_cols[i - 1] > 2:
                    sec_idx += 1
                col_to_sec[all_cols[i]] = sec_idx

        def sort_key(d):
            r, c = parse_pos(d.posicion)
            return (col_to_sec.get(c, 999), r, c)

        ordered = sorted(period_items, key=sort_key)

        self.save_state()
        for i, d in enumerate(ordered):
            d.codigo = f"{base}{start_num + i:0{width}d}{suffix}"

        self.items_by_codigo = {d.codigo: d for d in self.items}
        self.refresh_list()
        self.update_duplicates()
        self.show_cell_details(self.table.currentRow(), self.table.currentColumn())

        QMessageBox.information(
            self, "Listo",
            f"Actualizados {len(ordered)} código(s) de {self.period_title(period)}."
        )

    def on_save_json(self):
        if not self.items:
            return
        path, _ = QFileDialog.getSaveFileName(self, "Guardar JSON", "", "JSON (*.json)", options=QFileDialog.DontUseNativeDialog)
        if not path:
            return
        try:
            if self.root_data is not None and isinstance(self.root_data, dict) and "datosAG" in self.root_data:
                self.apply_global_ids_to_root()
                out = copy.deepcopy(self.root_data)
                out["datosAG"] = self._rebuild_datosAG()
                data = out
            else:
                data = self._sorted_items_list()
            with open(path, "w", encoding="utf-8") as f:
                json.dump(data, f, ensure_ascii=False, indent=2)
        except Exception as e:
            QMessageBox.critical(self, "Error", str(e))

    def _rebuild_datosAG(self) -> list:
        """Rebuild datosAG as list-of-lists grouped by id_form.

        Ordenamiento: (sección_de_columna, fila, columna).
        Una sección nueva comienza cuando el salto entre columnas consecutivas es > 2.
        Ej: cols {2,3,4} = sección 0, cols {7,8,9} = sección 1.
        Esto evita que bloques de columnas distintos se mezclen en el JSON.
        """
        groups: Dict[int, list] = {}
        for d in self.items:
            groups.setdefault(d.id_form, []).append(d)

        # Orden de sub-arrays: períodos configurados primero, resto por id_form
        seen: set = set()
        id_order = []
        for p in ["dia", "semana", "mes", "anio"]:
            pid = self.global_ids.get(p)
            if pid is not None and pid in groups and pid not in seen:
                id_order.append(pid)
                seen.add(pid)
        for pid in sorted(groups.keys()):
            if pid not in seen:
                id_order.append(pid)

        datosAG = []
        for pid in id_order:
            grp = groups[pid]

            # Detectar secciones de columnas: salto > 2 = nueva sección
            all_cols = sorted(set(parse_pos(d.posicion)[1] for d in grp))
            col_to_sec: Dict[int, int] = {}
            if all_cols:
                sec_idx = 0
                col_to_sec[all_cols[0]] = 0
                for i in range(1, len(all_cols)):
                    if all_cols[i] - all_cols[i - 1] > 2:
                        sec_idx += 1
                    col_to_sec[all_cols[i]] = sec_idx

            def _sort_key(d: CellItem, _m: Dict[int, int] = col_to_sec) -> tuple:
                r, c = parse_pos(d.posicion)
                return (_m.get(c, 999), r, c)

            sorted_grp = sorted(grp, key=_sort_key)
            datosAG.append([
                {"id_form": d.id_form, "label": d.label, "codigo": d.codigo,
                 "tipo": d.tipo, "deci": d.deci, "posicion": d.posicion, "valor": d.valor}
                for d in sorted_grp
            ])
        return datosAG

    def _sorted_items_list(self) -> list:
        """Flat sorted list of items used when root_data has no datosAG."""
        def sort_key(d):
            r, c = parse_pos(d.posicion)
            return (d.id_form, r, c)
        return [d.__dict__ for d in sorted(self.items, key=sort_key)]

    def _extract_client_code(self) -> Optional[str]:
        """Extract client suffix from existing codes (e.g., 'LUKE' from 'CD_00001LUKE')."""
        for d in self.items:
            m = re.match(r'^(?:CD|CS|CM|CA)_\d+([A-Z]+)$', d.codigo.strip().upper())
            if m:
                return m.group(1)
        return None

    def _next_codigo(self, period: str) -> str:
        """Generate next available codigo for a period."""
        prefix_map = {"dia": "CD", "semana": "CS", "mes": "CM", "anio": "CA"}
        prefix = prefix_map.get(period, "CD")
        client = self._extract_client_code() or ""
        max_seq = 0
        for d in self.items:
            m = re.match(rf'^{prefix}_[A-Z]*(\d+)', d.codigo.strip().upper())
            if m:
                max_seq = max(max_seq, int(m.group(1)))
        return f"{prefix}_{max_seq + 1:05d}{client}"

    def _swap_prefix(self, code: str, target: str) -> str:
        """Reusa el código del origen cambiando el prefijo (CD->CS->CM->CA).

        Ej: 'CD_MAS000001' copiado a semana queda 'CS_MAS000001'.
        Si el código no tiene un prefijo reconocible, genera uno nuevo."""
        prefix_map = {"dia": "CD", "semana": "CS", "mes": "CM", "anio": "CA"}
        target_prefix = prefix_map.get(target, "CD")
        code = (code or "").strip().upper()
        m = re.match(r'^(CD|CS|CM|CA)(.*)$', code)
        if m:
            return target_prefix + m.group(2)
        return self._next_codigo(target)

    def _validate_and_highlight_codigo(self, code: str):
        """Validate codigo format and highlight det_codigo if invalid."""
        valid = re.match(r'^(CD|CS|CM|CA)_([A-Z]+)?(\d+)([A-Z]*)$', code.strip().upper())
        if not valid:
            self.det_codigo.setStyleSheet("background:#CC0000; color:#fff; border:0; padding:8px;")
            self.det_codigo.setToolTip("Formato esperado: {CD|CS|CM|CA}_{texto}{numero}{cliente}")
            return
        suffix = valid.group(4)
        existing = self._extract_client_code()
        if existing and suffix and suffix != existing:
            self.det_codigo.setStyleSheet("background:#CC8800; color:#fff; border:0; padding:8px;")
            self.det_codigo.setToolTip(f"Sufijo '{suffix}' no coincide con el cliente existente '{existing}'")
        else:
            self.det_codigo.setStyleSheet("background:#2D2D44; color:#fff; border:0; padding:8px;")
            self.det_codigo.setToolTip("")

    def on_suggest_codigo(self):
        """Fill det_codigo with next available code for current period."""
        suggested = self._next_codigo(self.current_period)
        self.det_codigo.setText(suggested)
        self._validate_and_highlight_codigo(suggested)

    def on_import_excel(self):
        """Import Excel structure for a new client."""
        try:
            import openpyxl
        except ImportError:
            QMessageBox.critical(self, "Error", "openpyxl no está instalado.\nEjecuta: pip install openpyxl")
            return

        path, _ = QFileDialog.getOpenFileName(self, "Importar Excel", "", "Excel (*.xlsx *.xlsm)")
        if not path:
            return

        dlg = QDialog(self)
        dlg.setWindowTitle("Configurar Nuevo Cliente")
        layout = QVBoxLayout(dlg)
        form = QFormLayout()
        client_inp = QLineEdit()
        client_inp.setPlaceholderText("ej: LUKE")
        sheet_inp = QLineEdit()
        sheet_inp.setText("1")
        id_dia_inp = QLineEdit()
        id_sem_inp = QLineEdit()
        id_mes_inp = QLineEdit()
        id_anio_inp = QLineEdit()
        form.addRow("Código cliente (sufijo):", client_inp)
        form.addRow("Hoja (nombre o número):", sheet_inp)
        form.addRow("ID form Día:", id_dia_inp)
        form.addRow("ID form Semana:", id_sem_inp)
        form.addRow("ID form Mes:", id_mes_inp)
        form.addRow("ID form Año:", id_anio_inp)
        layout.addLayout(form)
        btns = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        btns.accepted.connect(dlg.accept)
        btns.rejected.connect(dlg.reject)
        layout.addWidget(btns)

        if dlg.exec() != QDialog.Accepted:
            return

        client_code = client_inp.text().strip().upper()
        if not client_code:
            QMessageBox.warning(self, "Error", "Debe ingresar un código de cliente.")
            return

        def to_int(s):
            try:
                return int(s.strip())
            except Exception:
                return None

        ids = {
            "dia": to_int(id_dia_inp.text()),
            "semana": to_int(id_sem_inp.text()),
            "mes": to_int(id_mes_inp.text()),
            "anio": to_int(id_anio_inp.text()),
        }
        configured_periods = [p for p in ["dia", "semana", "mes", "anio"] if ids.get(p) is not None]
        if not configured_periods:
            QMessageBox.warning(self, "Error", "Debe configurar al menos un ID de formulario.")
            return

        try:
            wb = openpyxl.load_workbook(path, data_only=True)
            sheet_str = sheet_inp.text().strip()
            ws = wb.worksheets[int(sheet_str) - 1] if sheet_str.isdigit() else wb[sheet_str]
        except Exception as e:
            QMessageBox.critical(self, "Error", f"No se pudo leer el Excel:\n{e}")
            return

        prefix_map = {"dia": "CD", "semana": "CS", "mes": "CM", "anio": "CA"}
        counters = {"dia": 1, "semana": 1, "mes": 1, "anio": 1}
        new_items = []
        for row in ws.iter_rows():
            for cell in row:
                val = cell.value
                if val is None or str(val).strip() == "":
                    continue
                label = str(val).strip()
                posicion = fmt_pos(cell.row - 1, cell.column - 1)
                for period in configured_periods:
                    pid = ids[period]
                    prefix = prefix_map[period]
                    seq = counters[period]
                    counters[period] += 1
                    new_items.append(CellItem(
                        id_form=pid,
                        label=label,
                        codigo=f"{prefix}_{seq:05d}{client_code}",
                        tipo=2,
                        deci=0,
                        posicion=posicion,
                        valor="",
                    ))

        if not new_items:
            QMessageBox.information(self, "Info", "No se encontraron datos en el Excel.")
            return

        ret = QMessageBox.question(
            self, "Confirmar",
            f"Se importarán {len(new_items)} campos para {len(configured_periods)} período(s).\n¿Reemplazar datos actuales?"
        )
        if ret != QMessageBox.Yes:
            return

        self.save_state()
        self.items = new_items
        self.items_by_codigo = {d.codigo: d for d in self.items}
        self.global_ids = ids
        self.pos_to_item = {}
        for d in self.items:
            p = self.get_period(d) or "dia"
            self.pos_to_item[(d.posicion, p)] = d

        type_map = {"dia": "d", "semana": "s", "mes": "m", "anio": "a"}
        cod_fechas = []
        for p in configured_periods:
            entry: Dict = {"id_form": ids[p], "tipo_val": type_map[p]}
            if p == "dia":
                entry["cod_fecha"] = f"FECHPROCE{client_code}"
            elif p == "mes":
                entry["cod_mes"] = f"MESCIERRE{client_code}"
                entry["cod_año"] = f"ANIOCIERRE{client_code}"
            elif p == "anio":
                entry["cod_año"] = f"ANIOCIERRE{client_code}PROCESO"
            cod_fechas.append(entry)

        self.root_data = {
            "bloquearFuentes": [],
            "validarDataForms": [],
            "formularioC": [{"fecha_cierre": f"FECHPROCE{client_code}", "cod_fechas": cod_fechas}],
            "datosAG": [],
        }

        self.build_groups()
        self.refresh_list()
        self.render_from_items()
        self.undo_stack.clear()
        self.redo_stack.clear()
        self.current_label.setText(f"Importados: {len(self.items)} items")
        if hasattr(self, "count_label") and self.count_label:
            self.count_label.setText(f"{len(self.items)} Items | Importados")
        QMessageBox.information(self, "Éxito", f"Se importaron {len(new_items)} campos del cliente {client_code}.")

    def refresh_list(self):
        if hasattr(self, "search_entry"):
            self.on_search_changed(self.search_entry.text())
        else:
            self.on_search_changed("")

    def render_from_items(self):
        # Ensure tabs exist for current items
        self.setup_tabs_from_items()
        # Calculate global grid size
        if not self.items:
            for tbl in self.tables.values():
                tbl.setRowCount(0)
                tbl.setColumnCount(0)
                tbl.clearContents()
            return
        max_r = 0
        max_c = 0
        for d in self.items:
            r, c = parse_pos(d.posicion)
            max_r = max(max_r, r)
            max_c = max(max_c, c)
        # Prepare each table
        for period, tbl in self.tables.items():
            tbl.setRowCount(max_r + 1)
            tbl.setColumnCount(max_c + 1)
            tbl.clearContents()
            for r in range(tbl.rowCount()):
                tbl.setRowHeight(r, 24)
            for c in range(tbl.columnCount()):
                tbl.setColumnWidth(c, 120)
            tbl.setHorizontalHeaderLabels([col_name(i) for i in range(tbl.columnCount())])
            tbl.setVerticalHeaderLabels([str(i + 1) for i in range(tbl.rowCount())])
        # Place items only into their period's table
        self.updating = True
        for d in self.items:
            self.place_item(d)
        self.update_duplicates()
        self.updating = False

    def place_item(self, d: CellItem):
        p = self.get_period(d) or "dia"
        tbl = self.tables.get(p)
        if not tbl:
            return
        r, c = parse_pos(d.posicion)
        qitem = QTableWidgetItem(d.label)
        qitem.setData(Qt.UserRole, d.codigo)
        # Aplicar color de duplicado si corresponde
        code_norm = (d.codigo or "").strip().upper()
        if code_norm and code_norm in self.duplicate_codes:
            qitem.setBackground(QBrush(QColor("#CC0000")))
            qitem.setForeground(QBrush(QColor("#FFFFFF")))
        tbl.setItem(r, c, qitem)

    def get_period(self, d: CellItem) -> Optional[str]:
        # 1° prioridad: prefijo del código (más confiable)
        code = d.codigo.strip().upper()
        if len(code) >= 2:
            prefix = code[:2]
            m = {"CD": "dia", "CS": "semana", "CM": "mes", "CA": "anio"}
            if prefix in m:
                return m[prefix]
        # 2° prioridad: coincidencia por id_form con global_ids
        for k in ["dia", "semana", "mes", "anio"]:
            val = self.global_ids.get(k)
            if val is not None and d.id_form == val:
                return k
        # Sin fallback por label — evita misclasificación de campos como
        # "INVENTARIO DIA/MES/AÑO" que contienen palabras de varios períodos
        return None

    def _active_periods(self) -> List[str]:
        """Períodos que realmente existen en el JSON cargado (detectados de formularioC + datosAG)."""
        if self.json_periods:
            return self.json_periods
        # Fallback si no hay JSON cargado: cualquier período con ID configurado
        return [p for p in ["dia", "semana", "mes", "anio"] if self.global_ids.get(p) is not None]

    def normalize_label(self, lbl: str) -> str:
        u = lbl.upper().strip()
        for k in [" DIA", " SEMANA", " MES", " AÑO", " ANIO"]:
            if u.endswith(k):
                u = u[: -len(k)]
                break
        return u

    def build_groups(self):
        self.groups = {}
        for d in self.items:
            p = self.get_period(d)
            base = self.normalize_label(d.label)
            if not p:
                p = "dia"
            if base not in self.groups:
                self.groups[base] = {}
            if p not in self.groups[base]:
                self.groups[base][p] = []
            self.groups[base][p].append(d)

    def move_group_for_item(self, pivot: CellItem, new_r: int, new_c: int):
        base = self.normalize_label(pivot.label)
        grp = self.groups.get(base, {})
        pr_r, pr_c = parse_pos(pivot.posicion)
        deltas: Dict[str, int] = {}
        
        for k, items_list in grp.items():
            if items_list:
                # Use the first item to determine the current column of this period block
                r, c = parse_pos(items_list[0].posicion)
                deltas[k] = c - pr_c
                
        max_c = 0
        max_r = 0
        # Determine current global grid size across items
        for d in self.items:
            rr, cc = parse_pos(d.posicion)
            max_r = max(max_r, rr)
            max_c = max(max_c, cc)
        max_c = max(max_c, (new_c + (max(deltas.values()) if deltas else 0)) + 1)
        max_r = max(max_r, new_r + 1)
        # Apply to all tables
        for tbl in self.tables.values():
            tbl.setRowCount(max_r)
            tbl.setColumnCount(max_c)
        
        # Bloquar señales de todas las tablas durante el movimiento de grupo
        for tbl in self.tables.values():
            tbl.blockSignals(True)
        try:
            for k, items_list in grp.items():
                # Clear old positions first
                for it in items_list:
                    old_r, old_c = parse_pos(it.posicion)
                    if old_r >= 0 and old_c >= 0:
                        p_it = self.get_period(it) or "dia"
                        tbl_it = self.tables.get(p_it)
                        if tbl_it:
                            tbl_it.setItem(old_r, old_c, QTableWidgetItem(""))

                # Update to new positions
                delta = deltas.get(k, 0)
                target_c = new_c + delta

                for it in items_list:
                    it.posicion = fmt_pos(new_r, target_c)
                    self.place_item(it)
        finally:
            for tbl in self.tables.values():
                tbl.blockSignals(False)

        self.pos_to_item = {}
        for d in self.items:
            p = self.get_period(d) or "dia"
            self.pos_to_item[(d.posicion, p)] = d

    def fill_id_fields(self, item: Optional[CellItem]):
        if not item:
            self.id_form_dia.setText("")
            self.id_form_semana.setText("")
            self.id_form_mes.setText("")
            self.id_form_anio.setText("")
            return
        base = self.normalize_label(item.label)
        grp = self.groups.get(base, {})
        
        def get_id(k):
            # Return ID from the first item in the list, or global fallback
            items_list = grp.get(k, [])
            if items_list:
                return str(items_list[0].id_form)
            val = self.global_ids.get(k)
            return str(val) if val else ""
            
        self.id_form_dia.setText(get_id("dia"))
        self.id_form_semana.setText(get_id("semana"))
        self.id_form_mes.setText(get_id("mes"))
        self.id_form_anio.setText(get_id("anio"))

    def on_update_ids(self):
        if not self.current_codigo:
            return
        ret = QMessageBox.question(self, "Confirmar", "¿Actualizar IDs del grupo?")
        if ret != QMessageBox.Yes:
            return
        self.save_state()
        item = self.items_by_codigo[self.current_codigo]
        base = self.normalize_label(item.label)
        grp = self.groups.get(base, {})
        def to_int(s: str) -> Optional[int]:
            s = s.strip()
            if not s:
                return None
            try:
                return int(s)
            except:
                return None
        v_d = to_int(self.id_form_dia.text())
        v_s = to_int(self.id_form_semana.text())
        v_m = to_int(self.id_form_mes.text())
        v_a = to_int(self.id_form_anio.text())
        
        if grp.get("dia") and v_d is not None:
            for it in grp["dia"]: it.id_form = v_d
        if grp.get("semana") and v_s is not None:
            for it in grp["semana"]: it.id_form = v_s
        if grp.get("mes") and v_m is not None:
            for it in grp["mes"]: it.id_form = v_m
        if grp.get("anio") and v_a is not None:
            for it in grp["anio"]: it.id_form = v_a
            
        self.global_ids = {
            "dia": v_d if v_d is not None else self.global_ids.get("dia"),
            "semana": v_s if v_s is not None else self.global_ids.get("semana"),
            "mes": v_m if v_m is not None else self.global_ids.get("mes"),
            "anio": v_a if v_a is not None else self.global_ids.get("anio"),
        }
        self.apply_global_ids_to_root()

    def extract_global_ids(self):
        ids = {"dia": None, "semana": None, "mes": None, "anio": None}
        detected: List[str] = []
        root = self.root_data

        # 1. Leer períodos desde formularioC
        try:
            if isinstance(root, dict) and isinstance(root.get("formularioC"), list):
                arr = root["formularioC"][0]
                cf = arr.get("cod_fechas", [])
                for e in cf:
                    tv = e.get("tipo_val")
                    if tv == "d":
                        ids["dia"] = e.get("id_form")
                        detected.append("dia")
                    elif tv == "s":
                        ids["semana"] = e.get("id_form")
                        detected.append("semana")
                    elif tv == "m":
                        ids["mes"] = e.get("id_form")
                        detected.append("mes")
                    elif tv == "a":
                        ids["anio"] = e.get("id_form")
                        detected.append("anio")
        except Exception:
            pass

        # 2. Validar: quitar períodos cuyo id_form no aparece realmente en datosAG
        if isinstance(root, dict) and isinstance(root.get("datosAG"), list):
            actual_ids: set = set()
            for group in root["datosAG"]:
                if isinstance(group, list):
                    for item in group:
                        if isinstance(item, dict) and "id_form" in item:
                            actual_ids.add(item["id_form"])
            for k in list(ids.keys()):
                if ids[k] is not None and ids[k] not in actual_ids:
                    ids[k] = None
                    detected = [p for p in detected if p != k]

            # 3. Fallback: si formularioC estaba vacío, detectar por prefijo de código en datosAG
            if not detected:
                prefix_map = {"CD": "dia", "CS": "semana", "CM": "mes", "CA": "anio"}
                id_to_period: Dict[int, str] = {}
                for group in root["datosAG"]:
                    if isinstance(group, list):
                        for item in group:
                            if isinstance(item, dict):
                                code = str(item.get("codigo", "")).strip().upper()
                                fid = item.get("id_form")
                                if len(code) >= 2 and code[:2] in prefix_map and fid is not None:
                                    p = prefix_map[code[:2]]
                                    if p not in detected:
                                        detected.append(p)
                                        ids[p] = fid

        self.global_ids = ids
        # Preservar orden canónico
        order = ["dia", "semana", "mes", "anio"]
        self.json_periods = [p for p in order if p in detected]

    def apply_global_ids_to_root(self):
        root = self.root_data
        try:
            if isinstance(root, dict) and isinstance(root.get("formularioC"), list):
                arr = root["formularioC"][0]
                cf = arr.get("cod_fechas", [])
                for e in cf:
                    tv = e.get("tipo_val")
                    if tv == "d" and self.global_ids.get("dia") is not None:
                        e["id_form"] = int(self.global_ids["dia"])
                    elif tv == "s" and self.global_ids.get("semana") is not None:
                        e["id_form"] = int(self.global_ids["semana"])
                    elif tv == "m" and self.global_ids.get("mes") is not None:
                        e["id_form"] = int(self.global_ids["mes"])
                    elif tv == "a" and self.global_ids.get("anio") is not None:
                        e["id_form"] = int(self.global_ids["anio"])
        except Exception:
            pass

    def update_root_with_items_and_ids(self, root):
        self.apply_global_ids_to_root()
        mapping = {d.codigo: d.posicion for d in self.items if d.codigo}
        def rec(v):
            if isinstance(v, dict):
                if "codigo" in v and "posicion" in v:
                    code = v.get("codigo")
                    if code in mapping:
                        v["posicion"] = mapping[code]
                for k in list(v.keys()):
                    v[k] = rec(v[k])
                return v
            elif isinstance(v, list):
                return [rec(e) for e in v]
            else:
                return v
        try:
            return rec(root)
        except Exception:
            return [d.__dict__ for d in self.items]

    def on_insert_row(self):
        idx = self.table.currentRow()
        col = self.table.currentColumn()
        if idx < 0:
            return
        self.save_state()
        updated = []
        for d in self.items:
            r, c = parse_pos(d.posicion)
            # Solo desplazar ítems del período activo; otros períodos tienen su propia grilla
            p = self.get_period(d) or "dia"
            if p == self.current_period and r >= idx:
                r += 1
            d.posicion = fmt_pos(r, c)
            updated.append(d)
        self.items = updated
        self.items_by_codigo = {d.codigo: d for d in self.items}
        
        self.pos_to_item = {}
        for d in self.items:
            p = self.get_period(d) or "dia"
            self.pos_to_item[(d.posicion, p)] = d
            
        self.render_from_items()
        
        # Restore selection to the same relative position (shifted down)
        # Note: idx was the row BEFORE insertion. The new empty row is at idx.
        # The content that was at idx is now at idx+1.
        # User usually wants to continue working near where they were.
        # Let's select the newly created empty row at the same column.
        
        # Force table update to ensure rows exist
        QApplication.processEvents()
        
        if self.table.rowCount() > idx:
            self.table.setCurrentCell(idx, col if col >= 0 else 0)
            self.table.setFocus()

    def on_insert_col(self):
        row = self.table.currentRow()
        idx = self.table.currentColumn()
        if idx < 0:
            return
        self.save_state()
        updated = []
        for d in self.items:
            r, c = parse_pos(d.posicion)
            # Solo desplazar ítems del período activo
            p = self.get_period(d) or "dia"
            if p == self.current_period and c >= idx:
                c += 1
            d.posicion = fmt_pos(r, c)
            updated.append(d)
        self.items = updated
        self.items_by_codigo = {d.codigo: d for d in self.items}
        
        self.pos_to_item = {}
        for d in self.items:
            p = self.get_period(d) or "dia"
            self.pos_to_item[(d.posicion, p)] = d
            
        self.render_from_items()
        
        # Restore selection to the newly created column
        QApplication.processEvents()
        
        if self.table.columnCount() > idx:
            self.table.setCurrentCell(row if row >= 0 else 0, idx)
            self.table.setFocus()

    def get_item_at(self, pos: str, period: str) -> Optional[CellItem]:
        return self.pos_to_item.get((pos, period))

    def _insert_item(self, new: CellItem):
        """Inserta un item nuevo en self.items en la posición correcta dentro de su grupo
        de id_form, respetando la sección de columna (salto > 2 = nueva sección)
        y el orden (fila, columna) dentro de esa sección."""
        new_r, new_c = parse_pos(new.posicion)

        # Detectar a qué sección pertenece la columna nueva
        # usando las columnas ya existentes del mismo id_form
        same_form = [d for d in self.items if d.id_form == new.id_form]
        all_cols = sorted(set(parse_pos(d.posicion)[1] for d in same_form))

        col_to_sec: Dict[int, int] = {}
        if all_cols:
            sec_idx = 0
            col_to_sec[all_cols[0]] = 0
            for i in range(1, len(all_cols)):
                if all_cols[i] - all_cols[i - 1] > 2:
                    sec_idx += 1
                col_to_sec[all_cols[i]] = sec_idx

        # Asignar sección al nuevo item: si su columna no existe aún,
        # inferirla buscando la columna más cercana sin salto > 2
        new_sec = 0
        if all_cols:
            # Buscar la sección que corresponde a new_c
            best_col = min(all_cols, key=lambda c: abs(c - new_c))
            if abs(best_col - new_c) <= 2:
                new_sec = col_to_sec[best_col]
            else:
                # Columna nueva forma su propia sección — determinar índice por posición
                new_sec = sum(1 for c in all_cols if c < new_c and all_cols[all_cols.index(c) + 1] - c > 2
                              if all_cols.index(c) + 1 < len(all_cols))

        # Buscar el último item del mismo id_form y misma sección con pos <= nueva
        insert_after = -1
        for i, d in enumerate(self.items):
            if d.id_form != new.id_form:
                continue
            dc_r, dc_c = parse_pos(d.posicion)
            d_sec = col_to_sec.get(dc_c, 999)
            if d_sec == new_sec and (dc_r, dc_c) <= (new_r, new_c):
                insert_after = i

        if insert_after == -1:
            # No hay item anterior en esa sección; insertar al inicio del grupo
            # o justo antes de la siguiente sección
            first_of_group = next((i for i, d in enumerate(self.items) if d.id_form == new.id_form), -1)
            if first_of_group == -1:
                self.items.append(new)
            else:
                self.items.insert(first_of_group, new)
        else:
            self.items.insert(insert_after + 1, new)

    def on_cell_changed(self, r: int, c: int):
        if self.updating:
            return
        itm = self.table.item(r, c)
        txt = itm.text().strip() if itm else ""
        pos = fmt_pos(r, c)
        
        # IMPROVED: Direct lookup with period awareness
        existing = self.pos_to_item.get((pos, self.current_period))
        
        changed = False
        if not txt:
            if existing:
                changed = True
        else:
            if not existing:
                changed = True
            elif existing.label != txt:
                changed = True
        
        if changed:
            self.save_state()
            
        if not txt:
            if existing:
                try:
                    self.items.remove(existing)
                except ValueError:
                    pass
                self.pos_to_item.pop((pos, self.current_period), None)
                self.items_by_codigo.pop(existing.codigo, None)
            return
        if existing:
            existing.label = txt
        else:
            # Determine appropriate ID for the current period
            current_id = 0
            if self.current_period in self.global_ids:
                val = self.global_ids[self.current_period]
                if val is not None:
                    current_id = val
                    
            new = CellItem(
                id_form=current_id,
                label=txt,
                codigo="",
                tipo=0,
                deci=0,
                posicion=pos,
                valor="",
            )
            self._insert_item(new)
            self.pos_to_item[(pos, self.current_period)] = new

        self.build_groups()
        self.show_cell_details(r, c)
        self.update_duplicates()
        self.refresh_list()

    def show_cell_details(self, r: int, c: int):
        pos = fmt_pos(r, c)
        
        # IMPROVED: Find item strictly for the current period
        it = self.pos_to_item.get((pos, self.current_period))

        if it:
            self.det_label.setText(it.label)
            self.det_codigo.setText(it.codigo)
            self.det_posicion.setText(it.posicion)
            self.det_id.setText(str(it.id_form))
            self.det_tipo.setText(str(it.tipo))
            self.det_deci.setText(str(it.deci))
            self.det_valor.setText(it.valor)
            if it.codigo:
                self._validate_and_highlight_codigo(it.codigo)
            else:
                self.det_codigo.setStyleSheet("background:#2D2D44; color:#fff; border:0; padding:8px;")
                self.det_codigo.setToolTip("")
        else:
            self.det_label.setText("")
            self.det_codigo.setText("")
            self.det_posicion.setText(f"{r}:{c}")
            self.det_id.setText("")
            self.det_tipo.setText("")
            self.det_deci.setText("")
            self.det_valor.setText("")
            self.det_codigo.setStyleSheet("background:#2D2D44; color:#fff; border:0; padding:8px;")
            self.det_codigo.setToolTip("")
            
    def on_detail_edited(self):
        if self.updating: return
        r = self.table.currentRow()
        c = self.table.currentColumn()
        if r < 0 or c < 0:
            return
        
        pos = fmt_pos(r, c)
        
        # IMPROVED: Find item strictly for the current period
        item = self.pos_to_item.get((pos, self.current_period))
        
        def safe_int(s):
            try: return int(s)
            except: return 0
            
        new_label = self.det_label.text()
        new_code = self.det_codigo.text().strip().upper() # Force upper and strip
        if new_code:
            self._validate_and_highlight_codigo(new_code)
        else:
            self.det_codigo.setStyleSheet("background:#2D2D44; color:#fff; border:0; padding:8px;")
            self.det_codigo.setToolTip("")
        new_id = safe_int(self.det_id.text())
        new_tipo = safe_int(self.det_tipo.text())
        new_deci = safe_int(self.det_deci.text())
        new_valor = self.det_valor.text()
        
        if not item:
            if not any([new_label, new_code, new_valor, new_id, new_tipo, new_deci]):
                return
            
            self.save_state()
            
            # If user didn't provide ID, try to guess from current period
            if new_id == 0:
                new_id = self.global_ids.get(self.current_period, 0)

            item = CellItem(
                id_form=new_id,
                label=new_label,
                codigo=new_code,
                tipo=new_tipo,
                deci=new_deci,
                posicion=pos,
                valor=new_valor
            )
            self._insert_item(item)
            self.pos_to_item[(pos, self.current_period)] = item

            if new_code:
                self.items_by_codigo[new_code] = item

            self.updating = True
            qitem = QTableWidgetItem(new_label)
            qitem.setData(Qt.UserRole, new_code)
            self.table.setItem(r, c, qitem)
            self.updating = False

            self.build_groups()
            self.refresh_list()
            self.current_label.setText(f"Cargados: {len(self.items)} items")
            if hasattr(self, "count_label") and self.count_label:
                self.count_label.setText(f"{len(self.items)} Items | Cargados")
            return

        if (item.label == new_label and
            item.codigo == new_code and
            item.id_form == new_id and
            item.tipo == new_tipo and
            item.deci == new_deci and
            item.valor == new_valor):
            return

        self.save_state()
        
        # DISABLE SIBLING SYNC - items are independent after creation
        # siblings = [] ...
        
        old_code = item.codigo
        if old_code != new_code:
            if old_code in self.items_by_codigo:
                del self.items_by_codigo[old_code]
            if new_code:
                self.items_by_codigo[new_code] = item
        
        # Update the main item
        item.label = new_label
        item.codigo = new_code
        item.id_form = new_id
        item.tipo = new_tipo
        item.deci = new_deci
        item.valor = new_valor
        
        # DISABLED SIBLING UPDATE LOOP
        # for sib in siblings: ...
        
        self.updating = True
        qitem = self.table.item(r, c)
        if not qitem:
            qitem = QTableWidgetItem(new_label)
            self.table.setItem(r, c, qitem)
        else:
            qitem.setText(new_label)
            qitem.setData(Qt.UserRole, new_code)
        self.updating = False
        
        self.build_groups() # Rebuild groups as label might have changed
        self.refresh_list()
        self.update_duplicates()

    def on_clear_all(self):
        self.save_state()
        self.items = []
        self.items_by_codigo = {}
        self.pos_to_item = {}
        self.groups = {}
        for tbl in self.tables.values():
            tbl.clearContents()
            tbl.setRowCount(0)
            tbl.setColumnCount(0)
        self.list.clear()
        self.current_label.setText("Cargados: 0 items")
        if hasattr(self, "count_label") and self.count_label:
            self.count_label.setText("0 Items | Cargados")

def main():
    app = QApplication(sys.argv)
    w = GridEditor()
    w.resize(1200, 700)
    w.show()
    sys.exit(app.exec())

if __name__ == "__main__":
    main()
